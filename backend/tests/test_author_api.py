from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app.config import get_settings
from backend.app.main import app

client = TestClient(app)


def test_ingest_local_author_samples_endpoint(
    tmp_path: Path,
    monkeypatch: object,
) -> None:
    db_path = tmp_path / "stylescribe.db"
    _set_db_path(monkeypatch, db_path)
    articles_dir = tmp_path / "articles"
    articles_dir.mkdir()
    _write_docx(articles_dir / "sample-title.docx", "API article text")
    metadata_path = tmp_path / "metadata.xlsx"
    _write_metadata(metadata_path)

    response = client.post(
        "/authors/ingest-local",
        json={
            "author_id": "v_vasanthi",
            "display_name": "V Vasanthi",
            "language": "ta",
            "articles_dir": str(articles_dir),
            "metadata_path": str(metadata_path),
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["author_id"] == "v_vasanthi"
    assert body["status"] == "completed"
    assert body["articles_seen"] == 1
    assert body["articles_ingested"] == 1
    assert body["articles_failed"] == 0
    assert body["metadata_rows_seen"] == 1
    assert body["warnings"] == []


def test_list_author_articles_endpoint(
    tmp_path: Path,
    monkeypatch: object,
) -> None:
    db_path = tmp_path / "stylescribe.db"
    _set_db_path(monkeypatch, db_path)
    articles_dir = tmp_path / "articles"
    articles_dir.mkdir()
    _write_docx(articles_dir / "sample-title.docx", "API article text")

    ingest_response = client.post(
        "/authors/ingest-local",
        json={
            "author_id": "v_vasanthi",
            "display_name": "V Vasanthi",
            "language": "ta",
            "articles_dir": str(articles_dir),
        },
    )
    assert ingest_response.status_code == 200

    response = client.get("/authors/v_vasanthi/articles")

    assert response.status_code == 200
    body = response.json()
    assert body == [
        {
            "article_id": "v_vasanthi:sample-title",
            "filename": "sample-title.docx",
            "title": None,
            "heading": None,
            "category": None,
            "text_char_count": len("API article text"),
            "url": None,
        }
    ]
    assert "extracted_text" not in body[0]


def _set_db_path(monkeypatch: object, db_path: Path) -> None:
    monkeypatch.setenv("STYLESCRIBE_DB_PATH", str(db_path))
    get_settings.cache_clear()


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _write_metadata(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Filename", "Title", "Category", "URL"])
    worksheet.append(
        [
            "sample-title.docx",
            "Sample Title",
            "Politics",
            "https://example.com/sample",
        ]
    )
    workbook.save(path)
