from pathlib import Path

from docx import Document
from openpyxl import Workbook

from backend.app.db.repository import StyleScribeRepository
from backend.app.services.author_ingestion_service import ingest_author_samples


def test_ingest_author_samples_with_metadata(tmp_path: Path) -> None:
    articles_dir = tmp_path / "articles"
    articles_dir.mkdir()
    _write_docx(articles_dir / "sample-title.docx", "Paragraph one")
    metadata_path = tmp_path / "metadata.xlsx"
    _write_metadata(metadata_path, filename="sample-title.docx")
    repository = StyleScribeRepository(tmp_path / "stylescribe.db")

    summary = ingest_author_samples(
        author_id="v_vasanthi",
        display_name="V Vasanthi",
        language="ta",
        articles_dir=articles_dir,
        metadata_path=metadata_path,
        repository=repository,
    )

    articles = repository.list_articles_for_author("v_vasanthi")
    assert summary.status == "completed"
    assert summary.articles_seen == 1
    assert summary.articles_ingested == 1
    assert summary.metadata_rows_seen == 1
    assert summary.warnings == []
    assert len(articles) == 1
    assert articles[0].title == "Sample Title"
    assert articles[0].category == "Politics"
    assert articles[0].url == "https://example.com/sample"


def test_ingest_author_samples_without_metadata(tmp_path: Path) -> None:
    articles_dir = tmp_path / "articles"
    articles_dir.mkdir()
    _write_docx(articles_dir / "article-one.docx", "Only article text")
    repository = StyleScribeRepository(tmp_path / "stylescribe.db")

    summary = ingest_author_samples(
        author_id="v_vasanthi",
        display_name="V Vasanthi",
        language="ta",
        articles_dir=articles_dir,
        metadata_path=None,
        repository=repository,
    )

    articles = repository.list_articles_for_author("v_vasanthi")
    assert summary.status == "completed"
    assert summary.metadata_rows_seen == 0
    assert summary.articles_ingested == 1
    assert articles[0].title is None
    assert articles[0].text_char_count == len("Only article text")


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _write_metadata(path: Path, filename: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Filename", "Title", "Heading", "Category", "URL"])
    worksheet.append(
        [
            filename,
            "Sample Title",
            "Sample Heading",
            "Politics",
            "https://example.com/sample",
        ]
    )
    workbook.save(path)
